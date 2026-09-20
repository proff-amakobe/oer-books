#!/usr/bin/env python3
"""Generate Blackboard Word guides from canonical weekly Markdown, never from a binary template."""
import argparse,json,re,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1];BASE=ROOT/'course-materials/blackboard';SRC=BASE/'source';OUT=BASE/'docx';QA=BASE/'qa'
NAVY='203864';BLUE='2F65A0';INK='202B38';LIGHT='EAF2FA';YELLOW='FFF4CC';GREEN='EAF3E5';WIDTH=10080

def xml(tag,**attrs):
 e=OxmlElement('w:'+tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 return e

def style(doc,name,size,color=INK,bold=False,before=0,after=6):
 s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(color);s.font.bold=bold
 s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after);s.paragraph_format.line_spacing=1.08
 s.paragraph_format.widow_control=True
 s.element.get_or_add_rPr().append(xml('rFonts',ascii='Arial',hAnsi='Arial',eastAsia='Arial',cs='Arial'))
 return s

def setup():
 d=Document();d.core_properties.author='Moody Amakobe';d.core_properties.title='Generative AI — Blackboard Module Guides';d.core_properties.subject='Graduate course | First Open Edition';d.core_properties.keywords='Generative AI, Blackboard, weekly module guides'
 style(d,'Normal',10.5);style(d,'Title',22,NAVY,True,0,8);style(d,'Subtitle',12,INK,False,0,8)
 for name,size,color in [('Heading 1',18,NAVY),('Heading 2',14,BLUE),('Heading 3',12,BLUE),('Heading 4',11,BLUE)]:
  s=style(d,name,size,color,True,8,6);s.paragraph_format.keep_with_next=True
 style(d,'Kicker',11,BLUE,True,0,5);style(d,'Section Band',10,NAVY,True,7,6)
 style(d,'Table Text',9.5,INK,False,0,3);style(d,'Guidance',9.5,INK,False,5,8);style(d,'Submission',10,INK,False,5,8)
 style(d,'Footer',8,NAVY,False,0,0);style(d,'List Bullet',10.5,INK,False,0,5);style(d,'List Number',10.5,INK,False,0,5)
 for name in ['Guidance','Submission']:d.styles[name].paragraph_format.keep_together=True
 # Real numbering definitions with explicit indent/tab geometry.
 root=d.part.numbering_part.element
 for n,fmt,char in [(70,'bullet','•'),(71,'decimal','%1.')]:
  a=xml('abstractNum',abstractNumId=n);a.append(xml('multiLevelType',val='singleLevel'));lv=xml('lvl',ilvl=0);lv.append(xml('start',val=1));lv.append(xml('numFmt',val=fmt));lv.append(xml('lvlText',val=char));lv.append(xml('suff',val='space'));lv.append(xml('lvlJc',val='left'))
  pp=xml('pPr');tabs=xml('tabs');tabs.append(xml('tab',val='num',pos=360));pp.append(tabs);pp.append(xml('ind',left=360,hanging=180));lv.append(pp);a.append(lv);root.append(a)
 d._bb_nextnum=80
 return d

def numid(d,ordered,start=1):
 n=d._bb_nextnum;d._bb_nextnum+=1;el=xml('num',numId=n);el.append(xml('abstractNumId',val=71 if ordered else 70));override=xml('lvlOverride',ilvl=0);override.append(xml('startOverride',val=start));el.append(override);d.part.numbering_part.element.append(el);return n

def inline(p,text):
 text=text.replace('\\$','$')
 pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)'
 for part in re.split(pattern,text):
  if not part:continue
  m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
  if m:
   h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True));run=xml('r');pr=xml('rPr');pr.append(xml('color',val=BLUE));pr.append(xml('u',val='single'));run.append(pr);t=xml('t');t.text=m[1];run.append(t);h.append(run);p._p.append(h)
  else:
   bold=part.startswith('**');italic=part.startswith('*') and not bold;code=part.startswith('`');clean=part[2:-2] if bold else part[1:-1] if italic or code else part
   run=p.add_run(clean);run.bold=bold;run.italic=italic
   if code:run.font.name='Courier New'

def para(d,text,sty='Normal'):
 p=d.add_paragraph(style=sty);inline(p,text);return p

def celltext(cell,text,header=False):
 p=cell.paragraphs[0];p.style='Table Text';inline(p,text)
 if header:
  for run in p.runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
 return p

def make_table(d,lines):
 rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines if not re.match(r'^\|\s*:?-',line)]
 n=len(rows[0]);widths=([450,1450,5900,2280] if rows[0][0]=='#' else [4450,2900,1730,1000]) if n==4 else ([1850,6450,1780] if rows[0][0]=='Pacing group' else [2550,6530,1000]) if n==3 else [WIDTH//n]*n
 if rows[0][0]=='Segment':widths=[1800,6880,1400]
 widths[-1]=WIDTH-sum(widths[:-1]);t=d.add_table(rows=0,cols=n);t.autofit=False
 pr=t._tbl.tblPr
 for name in ['tblW','tblInd','tblLayout','tblCellMar']:
  old=pr.find(qn('w:'+name))
  if old is not None:pr.remove(old)
 pr.append(xml('tblW',w=WIDTH,type='dxa'));pr.append(xml('tblInd',w=110,type='dxa'));pr.append(xml('tblLayout',type='fixed'))
 margins=xml('tblCellMar')
 for name,value in [('top',75),('bottom',75),('start',110),('end',110)]:margins.append(xml(name,w=value,type='dxa'))
 pr.append(margins);borders=xml('tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:borders.append(xml(side,val='single',sz=4,color='D4DFEA'))
 pr.append(borders)
 grid=t._tbl.tblGrid
 for e in list(grid):grid.remove(e)
 for w in widths:grid.append(xml('gridCol',w=w))
 for i,vals in enumerate(rows):
  cells=t.add_row().cells;rp=t.rows[-1]._tr.get_or_add_trPr();rp.append(xml('cantSplit'))
  if i==0:rp.append(xml('tblHeader'))
  for j,(c,txt) in enumerate(zip(cells,vals)):
   c.width=Inches(widths[j]/1440);tcpr=c._tc.get_or_add_tcPr();tcpr.find(qn('w:tcW')).set(qn('w:w'),str(widths[j]));tcpr.append(xml('vAlign',val='top'));tcpr.append(xml('shd',fill=NAVY if i==0 else LIGHT if i%2 else 'F5F8FB'));celltext(c,txt,i==0)
 return t

def section(d,meta,first=False):
 sec=d.sections[0] if first else d.add_section(WD_SECTION_START.NEW_PAGE)
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.75)
 sec.header_distance=Inches(.3);sec.footer_distance=Inches(.32)
 sec.footer.is_linked_to_previous=False;sec.header.is_linked_to_previous=False
 footer=sec.footer.paragraphs[0];footer.style='Footer';footer.text=f'Generative AI | First Open Edition | Week {meta["week"]} Module Guide'
 p=sec.footer.add_paragraph(meta['title'] if meta['week']<16 else 'Final Project Demonstration & Portfolio Synthesis',style='Footer')
 sec._sectPr.append(xml('pgNumType',start=1))
 return sec

def render_markdown(d,text):
 lines=text.splitlines();i=0;active_num=None;active_order=None
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;active_num=None;continue
  if line.startswith('<!-- metadata:'):i+=1;continue
  if line=='<!-- pagebreak -->':d.add_page_break();i+=1;continue
  if line.startswith('|'):
   table=[]
   while i<len(lines) and lines[i].strip().startswith('|'):table.append(lines[i]);i+=1
   make_table(d,table);continue
  m=re.match(r'^(#{1,4}) (.+)$',line)
  if m:
   text=m[2];level=len(m[1]);sty='Kicker' if text=='GENERATIVE AI' else 'Title' if re.match(r'Week \d+ Module Guide',text) else 'Section Band' if text.startswith('SECTION ') else 'Heading '+str(min(level,4));para(d,text,sty);i+=1;continue
  if line.startswith('> ['):
   m=re.match(r'> \[(GUIDANCE|SUBMISSION)\] (.+)',line);kind,txt=m.groups();p=para(d,txt,'Guidance' if kind=='GUIDANCE' else 'Submission');p.paragraph_format.left_indent=Inches(.1);p.paragraph_format.right_indent=Inches(.1);p._p.get_or_add_pPr().append(xml('shd',fill=YELLOW if kind=='GUIDANCE' else GREEN));i+=1;continue
  m=re.match(r'^(\d+\.|-) (.+)$',line)
  if m:
   ordered=m[1]!='-'
   if active_num is None or active_order!=ordered:active_num=numid(d,ordered,int(m[1][:-1]) if ordered else 1);active_order=ordered
   p=para(d,m[2],'List Number' if ordered else 'List Bullet');np=xml('numPr');np.append(xml('ilvl',val=0));np.append(xml('numId',val=active_num));p._p.get_or_add_pPr().append(np);i+=1;continue
  chunks=[line];i+=1
  while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||>|<!--|\d+\. |\- )',lines[i].strip()):chunks.append(lines[i].strip());i+=1
  para(d,' '.join(chunks))

def validate(path):
 with zipfile.ZipFile(path) as z:
  assert z.testzip() is None
  for n in z.namelist():
   if n.endswith(('.xml','.rels')):ET.fromstring(z.read(n))
 d=Document(path);assert len(d.paragraphs)>20
 # Fixed table geometry and widths; no merged cells or hidden text.
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 for t in d.tables:
  widths=[int(c.get(qn('w:w'))) for c in t._tbl.tblGrid];assert sum(widths)==WIDTH
  for row in t.rows:
   assert len(row.cells)==len(widths)
   for j,c in enumerate(row.cells):assert int(c._tc.tcPr.find(qn('w:tcW')).get(qn('w:w')))==widths[j]
 assert not d.element.xpath('.//w:vanish')
 return {'file':path.name,'zip':'PASS','xml':'PASS','reopen':'PASS','table_geometry':'PASS'}

def main():
 OUT.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True);combined=setup();results=[]
 for w in range(1,17):
  text=(SRC/f'week-{w:02}.md').read_text();meta=json.loads(re.search(r'<!-- metadata: (.+) -->',text)[1]);d=setup();section(d,meta,True);render_markdown(d,text);p=OUT/f'Generative_AI_Week{w:02}_Module_Guide.docx';d.save(p);results.append(validate(p));section(combined,meta,w==1);render_markdown(combined,text)
 p=OUT/'Generative_AI_Weeks01-16_Module_Guides.docx';combined.save(p);results.append(validate(p));(QA/'docx-validation.json').write_text(json.dumps(results,indent=2)+'\n');print('PASS: 16 individual guides and combined faculty guide validated.')
if __name__=='__main__':main()
